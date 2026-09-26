"""Kullanıcının istediği yapıya uyan gıda haberleri bülteni docx'i üretir:
banner+logo -> Türkiye (Hububat/Meyve-Sebze) -> Global Piyasa (Hububat/Meyve-Sebze)
-> TÜRİB Fiyatları -> TMO Fiyatları -> Kaynakça+Credit. Tüm gövde metni Times New Roman 12pt."""
from datetime import datetime
from pathlib import Path

import requests
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from gorseller import banner_uret, fiyat_hareketi_grafigi

RENK_KOYU = RGBColor(0x0F, 0x20, 0x40)
RENK_ORTA = RGBColor(0x1E, 0x4A, 0x82)
BEYAZ = RGBColor(0xFF, 0xFF, 0xFF)
RENK_KOYU_HEX = "0F2040"
RENK_ORTA_HEX = "1E4A82"

LOGO_YOLU = Path(__file__).parent / "logo.png"

# Bu kaynaklar Türkiye'ye özel kabul edilir; geri kalan her şey (Food Business News,
# Reddit vb.) "Dünya/Global Piyasa" bölümüne düşer.
TURKIYE_KAYNAK_ANAHTARLARI = ("dünya gazetesi", "linkedin/bloomberg", "reddit r/turkey")

UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"


def _turkiye_kaynagi_mi(kaynak):
    return any(a in kaynak.lower() for a in TURKIYE_KAYNAK_ANAHTARLARI)


def _hucre_arkaplan(hucre, hex_renk):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_renk)
    hucre._tc.get_or_add_tcPr().append(shd)


FONT_ADI = "Times New Roman"
GOVDE_PT = 12


def _paragraf(doc, metin, pt=GOVDE_PT, bold=False, color=None, align=None, italic=False):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    r = p.add_run(metin)
    r.font.name = FONT_ADI
    r.font.size = Pt(pt)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def _baslik_seridi(doc, baslik, tarih_str):
    tablo = doc.add_table(rows=1, cols=2)
    tablo.alignment = WD_TABLE_ALIGNMENT.CENTER
    tablo.autofit = False
    tablo.columns[0].width = Cm(13)
    tablo.columns[1].width = Cm(4)
    sol, sag = tablo.rows[0].cells
    _hucre_arkaplan(sol, RENK_KOYU_HEX)
    _hucre_arkaplan(sag, RENK_ORTA_HEX)
    for c in (sol, sag):
        c.vertical_alignment = 1

    p = sol.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(baslik)
    r.font.name = FONT_ADI
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = BEYAZ

    p2 = sag.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(tarih_str)
    r2.font.name = FONT_ADI
    r2.font.size = Pt(GOVDE_PT)
    r2.bold = True
    r2.font.color.rgb = BEYAZ
    doc.add_paragraph()


def _bolum_baslik(doc, no, metin):
    p = doc.add_paragraph()
    r = p.add_run(f"{no}. {metin}")
    r.font.name = FONT_ADI
    r.font.size = Pt(GOVDE_PT)
    r.bold = True
    r.font.color.rgb = RENK_KOYU
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)


def _alt_baslik(doc, metin):
    p = doc.add_paragraph()
    r = p.add_run(metin)
    r.font.name = FONT_ADI
    r.bold = True
    r.italic = True
    r.font.size = Pt(GOVDE_PT)
    r.font.color.rgb = RENK_KOYU
    p.paragraph_format.space_before = Pt(8)


def _kaynak_notu(doc, metin):
    _paragraf(doc, f"Kaynak: {metin}", italic=True, color=RGBColor(0x60, 0x60, 0x60))


def _footer_ekle(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


_gorsel_indirme_cache = {}


def _gorsel_indir(url, hedef_klasor, dosya_adi):
    if not url:
        return None
    if url in _gorsel_indirme_cache:
        return _gorsel_indirme_cache[url]
    try:
        r = requests.get(url, headers={"User-Agent": UA}, timeout=15)
        r.raise_for_status()
        yol = hedef_klasor / dosya_adi
        yol.write_bytes(r.content)
        _gorsel_indirme_cache[url] = yol
        return yol
    except Exception:
        _gorsel_indirme_cache[url] = None
        return None


def _fiyat_tablosu(doc, satirlar):
    tablo = doc.add_table(rows=1, cols=5)
    basliklar = ["Kaynak", "Ürün", "Önceki", "Bugün", "Değişim"]
    for i, b in enumerate(basliklar):
        c = tablo.rows[0].cells[i]
        _hucre_arkaplan(c, RENK_ORTA_HEX)
        p = c.paragraphs[0]
        r = p.add_run(b)
        r.font.name = FONT_ADI
        r.bold = True
        r.font.size = Pt(GOVDE_PT)
        r.font.color.rgb = BEYAZ
    for row in satirlar:
        cells = tablo.add_row().cells
        degisim = "-" if row["degisim_yuzde"] is None else f"{row['degisim_yuzde']:+.2f}%"
        degerler = [
            row["kaynak"], row["urun"],
            f"{row['ort_fiyat_onceki']}" if row["ort_fiyat_onceki"] is not None else "-",
            f"{row['ort_fiyat_son']} {row['birim'] or ''}",
            degisim,
        ]
        for i, v in enumerate(degerler):
            p = cells[i].paragraphs[0]
            r = p.add_run(str(v))
            r.font.name = FONT_ADI
            r.font.size = Pt(GOVDE_PT)


def bulten_olustur(cikti_yolu, tarih, haberler, tmo_turib_ozet):
    """tmo_turib_ozet: tmo_ozet.fiyat_ozeti_getir() çıktısı (tüm kaynaklar karışık) —
    burada TÜRİB (TURIB_ENDEKS, TURIB_NORMAL_SEANS) ve TMO kaynaklarına ayrıştırılır;
    diğer borsalar (Konya, ETB, Bandırma vb.) bu bültenin kapsamı dışında tutulur."""
    cikti_yolu = Path(cikti_yolu)
    calisma_klasoru = cikti_yolu.parent
    tarih_str = tarih.strftime("%d-%m-%Y")

    turib_ozet = [r for r in tmo_turib_ozet if r["kaynak"].startswith("TURIB")]
    tmo_ozet = [r for r in tmo_turib_ozet if r["kaynak"] == "TMO"]

    banner_yolu = calisma_klasoru / "_banner.png"
    grafik_turib_yolu = calisma_klasoru / "_grafik_turib.png"
    grafik_tmo_yolu = calisma_klasoru / "_grafik_tmo.png"
    banner_uret(banner_yolu)
    fiyat_hareketi_grafigi(grafik_turib_yolu, turib_ozet)
    fiyat_hareketi_grafigi(grafik_tmo_yolu, tmo_ozet)

    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    if LOGO_YOLU.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(str(LOGO_YOLU), width=Cm(2.2))
    doc.add_picture(str(banner_yolu), width=Cm(18))
    _baslik_seridi(doc, "Günlük Gıda ve Tarım Haberleri Bülteni", tarih_str)

    def _grupla(liste):
        sira, gruplu = [], {}
        for h in liste:
            if h["kaynak"] not in gruplu:
                gruplu[h["kaynak"]] = []
                sira.append(h["kaynak"])
            gruplu[h["kaynak"]].append(h)
        return sira, gruplu

    def _kaynak_yaz(kaynak_sirasi, gruplu):
        for kaynak in kaynak_sirasi:
            p = doc.add_paragraph()
            r = p.add_run(kaynak)
            r.font.name = FONT_ADI
            r.bold = True
            r.font.size = Pt(GOVDE_PT)
            r.font.color.rgb = RENK_ORTA
            for idx, h in enumerate(gruplu[kaynak][:5]):
                if h.get("gorsel_url"):
                    dosya_adi = f"_img_{abs(hash((h['kaynak'], h['baslik'])))}.jpg"
                    gorsel_yolu = _gorsel_indir(h["gorsel_url"], calisma_klasoru, dosya_adi)
                    if gorsel_yolu:
                        try:
                            doc.add_picture(str(gorsel_yolu), width=Cm(4))
                        except Exception:
                            pass
                _paragraf(doc, f"• {h['baslik']}")
                if h["ozet"]:
                    _paragraf(doc, f"   {h['ozet'][:260]}", color=RGBColor(0x40, 0x40, 0x40))

    def _bolge_yaz(liste, bos_mesaj):
        if not liste:
            _paragraf(doc, bos_mesaj)
            return
        hububat = [h for h in liste if h["kategori"] != "meyve_sebze"]
        meyve_sebze = [h for h in liste if h["kategori"] == "meyve_sebze"]
        if hububat:
            _alt_baslik(doc, "Hububat / Genel Tarım-Gıda")
            sira, gruplu = _grupla(hububat)
            _kaynak_yaz(sira, gruplu)
        if meyve_sebze:
            _alt_baslik(doc, "Meyve / Sebze")
            sira, gruplu = _grupla(meyve_sebze)
            _kaynak_yaz(sira, gruplu)

    turkiye_haberler = [h for h in haberler if _turkiye_kaynagi_mi(h["kaynak"])]
    global_haberler = [h for h in haberler if not _turkiye_kaynagi_mi(h["kaynak"])]

    # 1. Türkiye (Hububat / Meyve-Sebze)
    _bolum_baslik(doc, 1, "Türkiye")
    _bolge_yaz(turkiye_haberler, "Bugün Türkiye kaynaklarından haber alınamadı.")

    # 2. Global Piyasa (Hububat / Meyve-Sebze)
    _bolum_baslik(doc, 2, "Global Piyasa")
    _bolge_yaz(global_haberler, "Bugün global kaynaklardan haber alınamadı.")

    # 3. TÜRİB Fiyatları
    _bolum_baslik(doc, 3, "TÜRİB Fiyatları")
    if not turib_ozet:
        _paragraf(doc, "TÜRİB fiyat verisi alınamadı.")
    else:
        _fiyat_tablosu(doc, [r for r in turib_ozet if not r["anomali"]])
        if grafik_turib_yolu.exists():
            doc.add_picture(str(grafik_turib_yolu), width=Cm(14))
            _kaynak_notu(doc, "generalgrainrepo — TÜRİB Endeks + Normal Seans")

    # 4. TMO Fiyatları
    _bolum_baslik(doc, 4, "TMO Fiyatları")
    if not tmo_ozet:
        _paragraf(doc, "TMO fiyat verisi alınamadı.")
    else:
        _fiyat_tablosu(doc, [r for r in tmo_ozet if not r["anomali"]])
        if grafik_tmo_yolu.exists():
            doc.add_picture(str(grafik_tmo_yolu), width=Cm(14))
            _kaynak_notu(doc, "generalgrainrepo — TMO günlük fiyat bülteni")

    anomaliler = [r for r in (turib_ozet + tmo_ozet) if r["anomali"]]
    if anomaliler:
        doc.add_paragraph()
        _paragraf(doc, f"[UYARI] {len(anomaliler)} üründe veri anomalisi tespit edildi, tablolara dahil edilmedi.", italic=True)

    # 5. Kaynakça + Credit
    _bolum_baslik(doc, 5, "Kaynakça")
    for h in haberler:
        if h["link"]:
            _paragraf(doc, f"• [{h['kaynak']}] {h['baslik']} — {h['link']}")
    doc.add_paragraph()
    _paragraf(doc, "Fiyat verisi: github.com/utkusaraslan-hue/generalgrainrepo", italic=True)
    _paragraf(doc, f"Hazırlayan: Utku Sarıaslan — Sarıaslan Ticaret  |  Üretim zamanı: {datetime.now().strftime('%d.%m.%Y %H:%M')}", italic=True)

    _footer_ekle(doc)
    doc.save(cikti_yolu)
    return cikti_yolu
