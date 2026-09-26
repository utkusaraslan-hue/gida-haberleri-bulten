"""Sadece mekanik veri toplama adımı (LLM gerektirmez): haberleri ve
TMO/TÜRİB fiyat özetini çekip ham-veri/<tarih>/ klasörüne ayrı dosyalar
halinde yazar (tek bir .json yerine klasör + haberler.json + fiyatlar.xlsx +
gorseller/ alt klasörü — göz atması daha kolay olsun diye). Paragraf yazımı
ve HTML/PDF üretimi bu script'in dışında, .claude/skills/gunluk-bulten/
altındaki skill talimatına göre canlı bir Claude oturumunda yapılır."""
import json
import os
from datetime import datetime
from pathlib import Path

import openpyxl
import requests
from dotenv import load_dotenv
from openpyxl.styles import Font

from fetch_haberler import tum_haberleri_topla
from tmo_ozet import fiyat_ozeti_getir, tmo_il_ilce_fiyatlari_getir

BASE_DIR = Path(__file__).parent
UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"


def _haberler_json_yaz(haberler, cikti_klasoru):
    yol = cikti_klasoru / "haberler.json"
    with open(yol, "w", encoding="utf-8") as f:
        json.dump(haberler, f, ensure_ascii=False, indent=2, default=str)
    return yol


def _sayfa_doldur(ws, basliklar, satirlar, satir_to_row):
    ws.append(basliklar)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in satirlar:
        ws.append(satir_to_row(r))
    for col in ws.columns:
        genislik = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(genislik + 2, 40)


def _fiyatlar_xlsx_yaz(fiyat_satirlari, tmo_il_ilce_satirlari, cikti_klasoru):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Fiyatlar (özet)"
    _sayfa_doldur(
        ws1,
        ["Kaynak", "Ürün", "Birim", "Tarih (önceki)", "Fiyat (önceki)",
         "Tarih (bugün)", "Fiyat (bugün)", "Değişim %", "Anomali"],
        fiyat_satirlari,
        lambda r: [r["kaynak"], r["urun"], r["birim"], r["tarih_onceki"], r["ort_fiyat_onceki"],
                   r["tarih_son"], r["ort_fiyat_son"], r["degisim_yuzde"], "EVET" if r["anomali"] else ""],
    )

    ws2 = wb.create_sheet("TMO İl-İlçe")
    _sayfa_doldur(
        ws2,
        ["İl", "İlçe", "Ürün", "Birim", "Tarih (önceki)", "Fiyat (önceki)",
         "Tarih (bugün)", "Fiyat (bugün)", "Değişim %", "Anomali"],
        tmo_il_ilce_satirlari,
        lambda r: [r["il"], r["ilce"], r["urun"], r["birim"], r["tarih_onceki"], r["ort_fiyat_onceki"],
                   r["tarih_son"], r["ort_fiyat_son"], r["degisim_yuzde"], "EVET" if r["anomali"] else ""],
    )

    yol = cikti_klasoru / "fiyatlar.xlsx"
    wb.save(yol)
    return yol


def _gorselleri_indir(haberler, cikti_klasoru):
    gorseller_klasoru = cikti_klasoru / "gorseller"
    gorseller_klasoru.mkdir(exist_ok=True)
    indirilenler = 0
    for i, h in enumerate(haberler):
        url = h.get("gorsel_url")
        if not url:
            continue
        try:
            r = requests.get(url, headers={"User-Agent": UA}, timeout=15)
            r.raise_for_status()
            uzanti = ".jpg"
            if "webp" in r.headers.get("content-type", ""):
                uzanti = ".webp"
            elif "png" in r.headers.get("content-type", ""):
                uzanti = ".png"
            dosya_adi = f"{i:02d}_{h['kaynak'].replace(' ', '_').replace('/', '-')}{uzanti}"
            yol = gorseller_klasoru / dosya_adi
            yol.write_bytes(r.content)
            h["gorsel_dosya"] = f"gorseller/{dosya_adi}"
            indirilenler += 1
        except Exception:
            pass
    return indirilenler


def veri_topla_ve_kaydet(tarih=None):
    tarih = tarih or datetime.now()
    tarih_klasor_adi = tarih.strftime("%d-%m-%Y")
    cikti_klasoru = BASE_DIR / "ham-veri" / tarih_klasor_adi
    cikti_klasoru.mkdir(parents=True, exist_ok=True)

    load_dotenv(BASE_DIR / ".env")
    browser_use_key = os.environ.get("BROWSER_USE_API_KEY")

    print("Haberler toplanıyor...")
    haberler = tum_haberleri_topla(api_key=browser_use_key)
    print(f"  -> {len(haberler)} haber toplandı")

    print("Haber görselleri indiriliyor...")
    indirilen = _gorselleri_indir(haberler, cikti_klasoru)
    print(f"  -> {indirilen} görsel indirildi")

    print("TMO/TÜRİB fiyat özeti çekiliyor...")
    tmo_turib_ozet = fiyat_ozeti_getir()
    print(f"  -> {len(tmo_turib_ozet)} satır")

    print("TMO il/ilçe bazlı fiyatlar çekiliyor...")
    tmo_il_ilce = tmo_il_ilce_fiyatlari_getir()
    print(f"  -> {len(tmo_il_ilce)} satır")

    haberler_yolu = _haberler_json_yaz(haberler, cikti_klasoru)
    fiyat_yolu = _fiyatlar_xlsx_yaz(tmo_turib_ozet, tmo_il_ilce, cikti_klasoru)

    # tarih.json: küçük bir meta dosyası, hangi tarih için üretildiğini belirtir
    with open(cikti_klasoru / "tarih.json", "w", encoding="utf-8") as f:
        json.dump({"tarih": tarih_klasor_adi}, f, ensure_ascii=False, indent=2)

    print(f"Klasör hazır: {cikti_klasoru}")
    print(f"  - {haberler_yolu.name}")
    print(f"  - {fiyat_yolu.name}")
    print(f"  - gorseller/ ({indirilen} dosya)")
    return cikti_klasoru


if __name__ == "__main__":
    veri_topla_ve_kaydet()
